#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""NotebookLM(본드) 파이프라인 — 발표팩 → 노트북 → 슬라이드 → output/
(인포그래픽·브리핑 문서는 기본값이 아니다. 필요할 때만 --artifacts 로 추가한다.)

`nlm` CLI(notebooklm-mcp-cli)를 감싼다. 로이드가 `skills/notebooklm-handoff` 에서 호출한다.

  python .claude/tools/nlm_pipeline.py check
  python .claude/tools/nlm_pipeline.py run    --id <작업ID> [옵션]
  python .claude/tools/nlm_pipeline.py resume --id <작업ID>
  python .claude/tools/nlm_pipeline.py promote --id <작업ID>

설계상 지키는 것 —
  * 업로드 대상은 발표팩(07, 요르 06B가 쓰고 로이드가 검수한다)과 명시된 자료뿐이다. `00`~`06` 는 **코드가 거부한다**
    (중간 파일을 올리면 NotebookLM이 REMOVE된 주장을 되살린다).
  * 생성은 비동기다. 호출은 시작만 시킨다. `run` 은 완료까지 폴링한다.
  * 검사를 통과하지 못한 것을 `최종/<ID>/` 로 올리지 않는다. 이관은 `promote` 로 분리돼 있다.
  * 실패를 성공으로 적지 않는다. 모든 단계 결과가 `_nlm_run.json` 에 그대로 남는다.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

# ── 상수 ────────────────────────────────────────────────────────────────────

ROOT = Path(__file__).resolve().parents[2]

# 업로드가 금지된 workspace 파일. 번호 하나가 곧 계약이다.
FORBIDDEN_PREFIX = re.compile(r"^(0[0-6])_")  # 06(세부·비주얼 기획, 요르)도 완성 산출물이 아니다

# 아티팩트 정의: key → (nlm 하위명령, studio status 의 type, 다운로드 하위명령, 확장자)
ARTIFACTS = {
    "report": ("report", "report", "report", ".md"),
    "infographic": ("infographic", "infographic", "infographic", ".png"),
    "slides": ("slides", "slide_deck", "slide-deck", ".pdf"),
}

DEFAULT_ARTIFACTS = ["slides"]  # 사용자 결정(2026-09-18): 기본은 슬라이드만. 나머지는 필요할 때 요청.

# `--no-progress` 는 스트리밍 다운로드에만 있다. `download report` 에 붙이면
# "No such option" 으로 죽는다 (E-016).
PROGRESS_FLAG_OK = {"infographic", "slides"}

# 제작 지시 — 발표팩에 `# 제작 지시` 절이 없을 때 쓰는 비상용 기본값.
# 원본은 `요르/비주얼_발표.md` §4다. 문구를 고칠 일이 있으면 거기를 고치고 여기는 요약만 맞춘다.
FALLBACK_INSTRUCTION = """첨부한 Presentation Pack을 기준으로 만들어 주세요. 문서형 PPT로 옮기지 마세요.
- 슬라이드마다 「화면」의 주인공을 화면의 절반 이상으로 크게 그리세요. 넓은 여백에 작은 글자나 작은 인물만 두지 마세요.
- 덱 전체를 Art Direction의 색·서체 한 가족·화풍 하나로 통일하고, 본문 장 제목은 같은 크기·같은 위치에 두세요.
- 「제목」「화면 문구」의 항목 하나가 한 줄입니다. 그 위치에서만 줄을 바꾸세요. 「제목」이 "없음"이면 제목을 만들지 마세요.
- 「출처」가 "없음"인 장에는 출처를 쓰지 마세요. 출처는 왼쪽 아래 작은 한 줄로, 「한정」은 관련 숫자 옆 한 줄로 두세요.
- 첫 장은 발표의 얼굴로, 마지막 장은 첫 장의 모티프를 받아 결론 한 문장으로 닫으세요.
- 텍스트를 카드·박스에 나눠 담지 마세요. 관계·순서·수치는 도식·연표·큰 숫자로 보여 주세요.
- 캐릭터는 묘사대로 모든 장에서 같은 모습으로 그리세요. 로고·"공식" 표기·실존 인물 사진풍 얼굴은 만들지 마세요.
- 「발표자 메모」「근거」「최종 편집」은 화면에 표시하지 마세요.
- 자료에 없는 사실·수치를 추가하거나 계산하지 말고, 슬라이드 순서와 메시지를 바꾸지 마세요."""

GROUNDING = "Use only uploaded sources. Do not invent statistics, quotes, names, or examples not in the sources."


# ── 유틸 ────────────────────────────────────────────────────────────────────


def now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


def nlm_bin() -> str:
    """`nlm` 실행 파일. NLM_BIN 이 이긴다."""
    env = os.environ.get("NLM_BIN")
    if env and Path(env).exists():
        return env
    found = shutil.which("nlm")
    if found:
        return found
    for cand in (
        Path.home() / ".local" / "bin" / "nlm.exe",
        Path.home() / ".local" / "bin" / "nlm",
    ):
        if cand.exists():
            return str(cand)
    raise SystemExit(
        "nlm 을 찾지 못했다. `uv tool install notebooklm-mcp-cli` 후 다시 실행하거나 "
        "NLM_BIN 환경변수로 경로를 지정한다."
    )


def run_nlm(args: list[str], timeout: int = 900, profile: str | None = None) -> tuple[int, str, str]:
    """nlm 호출. (종료코드, stdout, stderr)."""
    cmd = [nlm_bin()] + args
    if profile:
        cmd += ["--profile", profile]
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    env["NO_COLOR"] = "1"
    env["TERM"] = "dumb"
    try:
        p = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            env=env,
            stdin=subprocess.DEVNULL,
        )
    except subprocess.TimeoutExpired:
        return 124, "", f"timeout after {timeout}s: {' '.join(args)}"
    return p.returncode, p.stdout or "", p.stderr or ""


def parse_json(stdout: str):
    """rich 장식이 섞여도 첫 JSON 값만 꺼낸다."""
    s = stdout.strip()
    if not s:
        return None
    for opener, closer in (("{", "}"), ("[", "]")):
        i = s.find(opener)
        j = s.rfind(closer)
        if i != -1 and j > i:
            try:
                return json.loads(s[i : j + 1])
            except json.JSONDecodeError:
                continue
    return None


def fail(state: dict, step: str, detail: str) -> None:
    state.setdefault("errors", []).append({"step": step, "detail": detail, "at": now()})
    log(f"  실패 · {step}: {detail}")


# ── 경로 ────────────────────────────────────────────────────────────────────


class Job:
    def __init__(self, job_id: str):
        self.id = job_id
        self.dir = ROOT / "작업" / job_id
        self.workspace = self.dir / "workspace"
        self.output = self.dir / "output"
        self.sources = self.dir / "sources"
        self.references = self.dir / "references"
        self.state_path = self.output / "_nlm_run.json"
        self.report_path = self.output / "_nlm_run.md"

    def ensure(self) -> None:
        self.output.mkdir(parents=True, exist_ok=True)

    def default_pack(self) -> Path | None:
        cands = sorted(self.workspace.glob("07_notebooklm_presentation_pack*.md"))
        return cands[-1] if cands else None

    def load_state(self) -> dict:
        if self.state_path.exists():
            return json.loads(self.state_path.read_text(encoding="utf-8"))
        return {}

    def save_state(self, state: dict) -> None:
        self.ensure()
        self.state_path.write_text(
            json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
        )


# ── 업로드 가드 ─────────────────────────────────────────────────────────────


def guard_source(path: Path) -> None:
    """검증 이전 단계의 중간 파일이 NotebookLM에 올라가는 것을 막는다."""
    if FORBIDDEN_PREFIX.match(path.name):
        raise SystemExit(
            f"업로드 거부: {path.name}\n"
            "  브리프·기획·조사 원본·검증 질답·세부 기획(00~06)은 NotebookLM에 올리지 않는다.\n"
            "  올리면 REMOVE로 폐기한 주장이 슬라이드에서 되살아난다 "
            "(skills/notebooklm-handoff §1)."
        )
    if not path.exists():
        raise SystemExit(f"업로드할 파일이 없다: {path}")


# ── 발표팩 파싱 ─────────────────────────────────────────────────────────────


def pack_instruction(pack_text: str) -> str:
    """팩의 `# 제작 지시` 절을 꺼낸다. 없으면 기본 지시."""
    m = re.search(r"^#\s*제작 지시.*?$(.*?)(?=^#\s|\Z)", pack_text, re.S | re.M)
    if m:
        body = m.group(1).strip()
        # 절차서 주석은 통째로 괄호로 묶여 있다 — 여러 줄에 걸쳐도 뗀다.
        body = re.sub(r"^\([^()]*?\)\s*$", "", body, flags=re.M).strip()
        body = re.sub(r"^```\w*\s*$", "", body, flags=re.M).strip()
        if len(body) > 40:
            return body
    return FALLBACK_INSTRUCTION


def pack_slide_count(pack_text: str) -> int | None:
    m = re.search(r"슬라이드 수\s*:\s*(\d+)", pack_text)
    if m:
        return int(m.group(1))
    slides = re.findall(r"^##\s*Slide\s+\d+", pack_text, re.M)
    return len(slides) or None


def focus_prompt(instruction: str, extra: str = "") -> str:
    """studio 생성용 focus 프롬프트. 근거 고정 문구를 항상 마지막에 붙인다."""
    parts = [instruction.strip()]
    if extra:
        parts.append(extra.strip())
    parts.append(GROUNDING)
    text = "\n".join(p for p in parts if p)
    return text[:9500]


# ── 숫자 대조 ───────────────────────────────────────────────────────────────

NUM_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")


def numbers_in(text: str) -> set[str]:
    out = set()
    for tok in NUM_RE.findall(text):
        norm = tok.replace(",", "")
        norm = norm.rstrip(".")
        if norm:
            out.add(norm)
    return out


HEADING_NUM_RE = re.compile(r"^\s*(?:#{1,6}\s*)?\d{1,2}(?:\.\d{1,2})*[.)]?\s", re.M)


def strip_structure(text: str) -> str:
    """절 번호·목록 번호를 뗀다. `## 2.1` 의 2.1 이 새 수치로 잡히던 오탐(2026-09-17 실측)."""
    return HEADING_NUM_RE.sub(" ", text)


def novel_numbers(pack_text: str, artifact_text: str) -> list[str]:
    """팩에 없는데 생성물에 나온 숫자. 절·장번호와 0~50 정수는 뺀다."""
    artifact_text = strip_structure(artifact_text)
    base = numbers_in(pack_text)
    found = []
    for n in sorted(numbers_in(artifact_text), key=lambda x: (-len(x), x)):
        if n in base:
            continue
        try:
            if "." not in n and 0 <= int(n) <= 50:
                continue
        except ValueError:
            pass
        found.append(n)
    return found[:40]


# ── 검사 ────────────────────────────────────────────────────────────────────


def inspect_pptx(path: Path) -> dict:
    """NotebookLM 덱은 슬라이드 한 장이 통 이미지다(E-017). 빈 장으로 오판하지 않는다."""
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = []
    text_all = []
    pictures = 0
    for i, s in enumerate(prs.slides, 1):
        t = " ".join(
            sh.text_frame.text for sh in s.shapes if getattr(sh, "has_text_frame", False)
        ).strip()
        pics = sum(
            1 for sh in s.shapes if str(getattr(sh, "shape_type", "")).startswith("PICTURE")
        )
        pictures += pics
        text_all.append(t)
        slides.append({"n": i, "chars": len(t), "pics": pics, "head": t[:60].replace("\n", " ")})

    image_only = pictures >= len(slides) > 0 and all(s["chars"] < 5 for s in slides)
    blank = [s["n"] for s in slides if s["chars"] < 5 and s["pics"] == 0]
    return {
        "ok": len(slides) > 0 and not blank,
        "slide_count": len(slides),
        "image_only": image_only,
        "blank_slides": blank,
        "long_slides": [s["n"] for s in slides if s["chars"] > 600],
        "slides": slides,
        "text": "\n".join(text_all),
        "text_checkable": not image_only,
    }


def inspect_pdf(path: Path) -> dict:
    import fitz

    d = fitz.open(str(path))
    pages = [d[i].get_text() for i in range(d.page_count)]
    images = [len(d[i].get_images()) for i in range(d.page_count)]
    d.close()
    image_only = bool(pages) and all(len(t.strip()) < 5 and n >= 1 for t, n in zip(pages, images))
    blank = [i + 1 for i, (t, n) in enumerate(zip(pages, images)) if len(t.strip()) < 5 and n == 0]
    return {
        "ok": len(pages) > 0 and not blank,
        "page_count": len(pages),
        "image_only": image_only,
        "blank_pages": blank,
        "text": "\n".join(pages),
        "text_checkable": not image_only,
    }


def inspect_png(path: Path) -> dict:
    from PIL import Image

    with Image.open(path) as im:
        w, h = im.size
    size = path.stat().st_size
    return {
        "ok": size > 20_000 and w >= 600 and h >= 400,
        "width": w,
        "height": h,
        "bytes": size,
        "text": "",
    }


def inspect_md(path: Path) -> dict:
    t = path.read_text(encoding="utf-8", errors="replace")
    return {"ok": len(t.strip()) > 200, "chars": len(t), "text": t}


def inspect_file(path: Path) -> dict:
    suf = path.suffix.lower()
    try:
        if suf == ".pptx":
            return inspect_pptx(path)
        if suf == ".pdf":
            return inspect_pdf(path)
        if suf in (".png", ".jpg", ".jpeg", ".webp"):
            return inspect_png(path)
        return inspect_md(path)
    except Exception as e:  # 검사 자체가 깨져도 파이프라인을 죽이지 않는다
        return {"ok": False, "error": f"{type(e).__name__}: {e}", "text": ""}


# ── 단계 ────────────────────────────────────────────────────────────────────


def ensure_auth(profile: str | None, log_fn=print) -> tuple[bool, str]:
    """로그인 확인 → 실패하면 저장된 프로필로 무인 갱신(`nlm auth refresh`)까지 시도한다.

    쿠키 만료는 매번 사람이 `nlm login`(브라우저)을 다시 열 일이 아니다 — 저장된 전용 크롬
    프로필이 있으면 로이드가 headless로 재발급받을 수 있다(E-019, 2026-09-18 실측).
    반환: (성공 여부, 마지막으로 확인한 원문 출력).
    """
    code, out, err = run_nlm(["login", "--check"], timeout=180, profile=profile)
    text = (out + err).strip()
    if code == 0:
        return True, text

    log_fn(text)
    log_fn("→ 인증 확인 실패. 저장된 세션으로 무인 갱신을 시도한다 (`nlm auth refresh`)...")
    rcode, rout, rerr = run_nlm(["auth", "refresh"], timeout=180, profile=profile)
    log_fn((rout + rerr).strip())
    if rcode != 0:
        return False, text

    code2, out2, err2 = run_nlm(["login", "--check"], timeout=180, profile=profile)
    if code2 == 0:
        return True, out2.strip()
    return False, (out2 + err2).strip()


def step_check(profile: str | None) -> int:
    try:
        binary = nlm_bin()
    except SystemExit as e:
        print(e)
        return 1
    print(f"nlm: {binary}")
    code, out, err = run_nlm(["--version"], timeout=120, profile=None)
    print((out or err).strip().splitlines()[0] if (out or err).strip() else "(버전 확인 실패)")

    ok, text = ensure_auth(profile)
    print(text)
    if not ok:
        print("\n→ 무인 갱신도 실패했다. 저장된 프로필 자체가 없거나 완전히 만료됐다는 뜻이다.")
        print("   사용자가 브라우저에서 다시 로그인해야 한다:")
        print("   nlm login")
        return 1
    print("\n인증 OK — 파이프라인을 돌릴 수 있다.")
    return 0


def step_create_notebook(job: Job, state: dict, title: str, profile: str | None) -> bool:
    code, out, err = run_nlm(["notebook", "create", title, "--json"], timeout=300, profile=profile)
    data = parse_json(out)
    if code != 0 or not data or not data.get("notebook_id"):
        fail(state, "notebook_create", (err or out).strip()[:500])
        return False
    state["notebook_id"] = data["notebook_id"]
    state["notebook_url"] = data.get("url", "")
    state["notebook_title"] = data.get("title", title)
    log(f"  노트북 생성: {state['notebook_id']}")
    return True


def step_persona(job: Job, state: dict, instruction: str, profile: str | None) -> None:
    """노트북 페르소나. Briefing Doc 처럼 프롬프트 슬롯이 없는 아티팩트에도 걸린다."""
    prompt = focus_prompt(instruction)
    code, out, err = run_nlm(
        ["chat", "configure", state["notebook_id"], "--goal", "custom", "--prompt", prompt],
        timeout=300,
        profile=profile,
    )
    state["persona_set"] = code == 0
    if code != 0:
        fail(state, "chat_configure", (err or out).strip()[:300])
    else:
        log("  노트북 페르소나 설정")


def step_upload(job: Job, state: dict, files: list[Path], urls: list[str], profile: str | None) -> bool:
    uploaded = state.setdefault("sources", [])
    ok_any = False
    for f in files:
        guard_source(f)
        log(f"  업로드: {f.name}")
        code, out, err = run_nlm(
            [
                "source", "add", state["notebook_id"],
                "--file", str(f),
                "--title", f.stem,
                "--wait", "--wait-timeout", "600",
                "--json",
            ],
            timeout=900,
            profile=profile,
        )
        try:
            ref = f.relative_to(ROOT).as_posix()
        except ValueError:
            ref = f.as_posix()
        entry = {"kind": "file", "ref": ref, "ok": code == 0}
        if code != 0:
            entry["error"] = (err or out).strip()[:300]
            fail(state, f"source_add:{f.name}", entry["error"])
        else:
            data = parse_json(out) or {}
            entry["source_id"] = data.get("id") or data.get("source_id")
            ok_any = True
        uploaded.append(entry)
    for u in urls:
        log(f"  업로드(URL): {u}")
        code, out, err = run_nlm(
            ["source", "add", state["notebook_id"], "--url", u, "--wait", "--json"],
            timeout=900,
            profile=profile,
        )
        entry = {"kind": "url", "ref": u, "ok": code == 0}
        if code != 0:
            entry["error"] = (err or out).strip()[:300]
            fail(state, f"source_add:{u}", entry["error"])
        else:
            ok_any = True
        uploaded.append(entry)
    return ok_any


def build_create_args(kind: str, state: dict, opts: argparse.Namespace, instruction: str) -> list[str]:
    nb = state["notebook_id"]
    lang = opts.language
    if kind == "report":
        args = ["report", "create", nb, "--format", opts.report_format]
        if opts.report_format == "Create Your Own":
            args += ["--prompt", focus_prompt(instruction)]
        if lang:
            args += ["--language", lang]
    elif kind == "infographic":
        args = [
            "infographic", "create", nb,
            "--orientation", opts.orientation,
            "--detail", opts.detail,
            "--style", opts.style,
            "--focus", focus_prompt(instruction, "핵심 수치와 근거 각주를 반드시 포함한다."),
        ]
        if lang:
            args += ["--language", lang]
    else:  # slides
        args = [
            "slides", "create", nb,
            "--format", opts.slide_format,
            "--length", opts.slide_length,
            "--focus", focus_prompt(instruction),
        ]
        if lang:
            args += ["--language", lang]
    return args + ["--confirm", "--json"]


def step_create_artifacts(
    job: Job, state: dict, kinds: list[str], opts: argparse.Namespace, instruction: str
) -> None:
    arts = state.setdefault("artifacts", {})
    for kind in kinds:
        if arts.get(kind, {}).get("artifact_id"):
            log(f"  {kind}: 이미 생성됨 — 건너뜀")
            continue
        args = build_create_args(kind, state, opts, instruction)
        log(f"  생성 시작: {kind}")
        code, out, err = run_nlm(args, timeout=900, profile=opts.profile)
        data = parse_json(out) or {}
        aid = data.get("artifact_id")
        if code != 0 or not aid:
            arts[kind] = {"artifact_id": None, "status": "create_failed"}
            fail(state, f"create:{kind}", (err or out).strip()[:400])
            continue
        arts[kind] = {
            "artifact_id": aid,
            "status": data.get("status", "in_progress"),
            "created_at": now(),
        }
        log(f"    artifact_id = {aid}")
        job.save_state(state)


def step_poll(job: Job, state: dict, kinds: list[str], opts: argparse.Namespace) -> None:
    """studio status 를 폴링한다. 생성은 비동기다 — 호출만으로 끝나지 않는다."""
    arts = state.get("artifacts", {})
    pending = {
        k: arts[k]["artifact_id"]
        for k in kinds
        if arts.get(k, {}).get("artifact_id") and arts[k].get("status") != "completed"
    }
    if not pending:
        return
    deadline = time.time() + opts.timeout
    interval = 20
    while pending and time.time() < deadline:
        code, out, err = run_nlm(
            ["studio", "status", state["notebook_id"], "--json"],
            timeout=300,
            profile=opts.profile,
        )
        rows = parse_json(out)
        # E-025(2026-09-18): run_in_background로 분리된 프로세스에서 이 호출이
        # 유효한 JSON을 stdout에 내면서도 code != 0 을 반복 반환하는 사례를 관찰했다.
        # 포그라운드에서는 같은 호출이 매번 code 0 이었다 — 원인은 특정하지 못했다(추정: 콘솔
        # 핸들이 없는 detached 프로세스에서 nlm CLI 내부의 종료 후 처리가 흔들리는 것으로 보임).
        # code 를 무조건 신뢰하면 이미 완성된 아티팩트를 영영 "실패"로 오판해 폴링만 반복하므로,
        # rows 가 유효한 리스트면 code 와 무관하게 그 내용을 신뢰한다. code != 0 은 진단용으로만 남긴다.
        if code != 0 and isinstance(rows, list):
            log(f"    studio_status: code={code} 인데 유효한 JSON을 받음 — 내용을 신뢰한다(E-025)")
        if isinstance(rows, list):
            by_id = {r.get("artifact_id") or r.get("id"): r for r in rows}
            for kind, aid in list(pending.items()):
                row = by_id.get(aid)
                if not row:
                    continue
                st = row.get("status", "")
                if st and st != arts[kind].get("status"):
                    arts[kind]["status"] = st
                    log(f"    {kind}: {st}")
                if st == "completed":
                    arts[kind]["completed_at"] = now()
                    pending.pop(kind)
                elif st == "failed":
                    fail(state, f"generate:{kind}", "NotebookLM이 생성에 실패했다(status=failed)")
                    pending.pop(kind)
        else:
            fail(state, "studio_status", (err or out).strip()[:300])
        job.save_state(state)
        if pending:
            time.sleep(interval)
            interval = min(interval + 10, 60)
    for kind in pending:
        arts[kind]["status"] = arts[kind].get("status", "in_progress")
        fail(
            state,
            f"timeout:{kind}",
            f"{opts.timeout}s 안에 완료되지 않았다. `resume` 으로 이어받는다.",
        )


def step_download(job: Job, state: dict, kinds: list[str], opts: argparse.Namespace) -> None:
    job.ensure()
    arts = state.get("artifacts", {})
    for kind in kinds:
        info = arts.get(kind, {})
        if info.get("status") != "completed" or not info.get("artifact_id"):
            continue
        _, _, dl_cmd, ext = ARTIFACTS[kind]
        targets = [(ext, [])]
        if kind == "slides":
            # 덱은 장마다 통 이미지라 pptx 로 받아도 글자를 못 고친다(E-017).
            # 기본은 PDF 하나. pptx 가 필요하면 --slides-pptx.
            targets = [(".pdf", ["--format", "pdf"])]
            if getattr(opts, "slides_pptx", False):
                targets.append((".pptx", ["--format", "pptx"]))
        files = info.setdefault("files", {})
        for suffix, extra in targets:
            dest = job.output / f"{job.id}_{kind}{suffix}"
            # 대상 파일이 이미 있으면 nlm 이 "Download failed" 로 죽는다 (E-016).
            # 받아 둔 것을 다시 받지 않는다 — `resume` 이 멀쩡한 파일을 깨뜨리지 않게.
            if dest.exists() and dest.stat().st_size > 0 and not getattr(opts, "redownload", False):
                files[suffix] = dest.relative_to(ROOT).as_posix()
                log(f"  이미 받음: {dest.name} — 건너뜀")
                continue
            if dest.exists() and getattr(opts, "redownload", False):
                dest.unlink()
            log(f"  다운로드: {dest.name}")
            cmd = ["download", dl_cmd, state["notebook_id"], "--id", info["artifact_id"],
                   "--output", str(dest)]
            if kind in PROGRESS_FLAG_OK:
                cmd.append("--no-progress")
            code, out, err = run_nlm(cmd + extra, timeout=1800, profile=opts.profile)
            if code != 0 or not dest.exists() or dest.stat().st_size == 0:
                fail(state, f"download:{kind}{suffix}", (err or out).strip()[:400])
                continue
            files[suffix] = dest.relative_to(ROOT).as_posix()
        job.save_state(state)


def step_inspect(job: Job, state: dict, pack_text: str) -> None:
    """기계 검사. 눈으로 볼 것 세 가지는 로이드가 따로 본다(skills §5)."""
    checks = state.setdefault("inspection", {})
    expected = pack_slide_count(pack_text)
    state["pack_slide_count"] = expected
    for kind, info in state.get("artifacts", {}).items():
        for suffix, rel in (info.get("files") or {}).items():
            p = ROOT / rel
            if not p.exists():
                continue
            res = inspect_file(p)
            text = res.pop("text", "")
            if res.get("text_checkable") is False:
                res["novel_numbers"] = []
                res["number_check"] = "불가 — 이미지 슬라이드다. 숫자는 눈으로 본다"
            else:
                res["novel_numbers"] = novel_numbers(pack_text, text) if text else []
            if kind == "slides" and suffix in (".pptx", ".pdf") and expected:
                n = res.get("slide_count") or res.get("page_count")
                res["expected_slides"] = expected
                # NotebookLM은 표지를 한 장 덧붙인다(2026-09-17 실측). +1 까지는 정상.
                res["slide_count_match"] = n in (expected, expected + 1)
                if n == expected + 1:
                    res["note"] = "팩보다 1장 많다 — 표지로 보인다. 첫 장을 확인한다"
            checks[f"{kind}{suffix}"] = res
            flag = "OK" if res.get("ok") else "확인 필요"
            log(f"  검사 {kind}{suffix}: {flag}"
                + (f" · 팩에 없는 숫자 {len(res['novel_numbers'])}건" if res["novel_numbers"] else ""))


def write_report(job: Job, state: dict) -> None:
    L = []
    L.append(f"# NotebookLM 실행 기록 — {job.id}")
    L.append("")
    L.append(f"- 실행 시각: {state.get('started_at')} → {state.get('finished_at')}")
    L.append(f"- 경로: CLI 자동 (`nlm` / notebooklm-mcp-cli)")
    L.append(f"- 노트북: {state.get('notebook_title','')} · `{state.get('notebook_id','')}`")
    if state.get("notebook_url"):
        L.append(f"- 노트북 URL: {state['notebook_url']}")
    L.append(f"- 팩 슬라이드 수: {state.get('pack_slide_count')}")
    L.append("")
    L.append("## 올린 자료")
    for s in state.get("sources", []):
        L.append(f"- [{'OK' if s.get('ok') else '실패'}] {s.get('kind')}: {s.get('ref')}")
    L.append("")
    L.append("## 생성물")
    L.append("")
    L.append("| 종류 | 상태 | artifact_id | 파일 |")
    L.append("| --- | --- | --- | --- |")
    for kind, info in state.get("artifacts", {}).items():
        files = ", ".join((info.get("files") or {}).values()) or "—"
        L.append(f"| {kind} | {info.get('status')} | `{info.get('artifact_id')}` | {files} |")
    L.append("")
    L.append("## 기계 검사")
    for key, res in state.get("inspection", {}).items():
        L.append(f"- **{key}** — {'OK' if res.get('ok') else '확인 필요'}")
        for k in ("slide_count", "expected_slides", "slide_count_match", "note",
                  "image_only", "number_check", "blank_slides", "long_slides",
                  "page_count", "blank_pages", "width", "height", "bytes",
                  "chars", "error"):
            if k in res and res[k] not in (None, [], ""):
                L.append(f"  - {k}: {res[k]}")
        if res.get("novel_numbers"):
            L.append(f"  - **팩에 없는 숫자**: {', '.join(res['novel_numbers'])}")
            L.append("    → 근거 없이 생긴 값일 수 있다. 로이드가 직접 확인한다.")
    if state.get("errors"):
        L.append("")
        L.append("## 실패·미완료")
        for e in state["errors"]:
            L.append(f"- `{e['step']}` — {e['detail']}")
    L.append("")
    L.append("## 다음")
    L.append("1. 전 장 렌더 후 눈으로 판정 (`skills/notebooklm-handoff` §4) — 사실(팩에 없는 주장 / "
             "Qualification 생존 / 검증 기록 노출) + Render Gate.")
    L.append("2. `skills/cross-check` 로 크게 박힌 수치를 원문과 대조.")
    L.append(f"3. 둘 다 닫히면 `python .claude/tools/nlm_pipeline.py promote --id {job.id}`.")
    L.append("")
    job.report_path.write_text("\n".join(L), encoding="utf-8")
    log(f"  기록: {job.report_path.relative_to(ROOT).as_posix()}")


# ── 명령 ────────────────────────────────────────────────────────────────────


def cmd_run(opts: argparse.Namespace) -> int:
    job = Job(opts.id)
    if not job.dir.exists():
        raise SystemExit(f"작업 폴더가 없다: {job.dir}")
    job.ensure()

    pack = Path(opts.pack) if opts.pack else job.default_pack()
    if not pack or not pack.exists():
        raise SystemExit(
            "발표팩을 찾지 못했다. "
            f"{job.workspace}/07_notebooklm_presentation_pack.md 를 먼저 만든다."
        )
    guard_source(pack)
    pack_text = pack.read_text(encoding="utf-8", errors="replace")
    instruction = pack_instruction(pack_text)

    extra_files = [Path(f) if Path(f).is_absolute() else ROOT / f for f in (opts.source or [])]
    for f in extra_files:
        guard_source(f)

    kinds = [k.strip() for k in opts.artifacts.split(",") if k.strip()]
    for k in kinds:
        if k not in ARTIFACTS:
            raise SystemExit(f"모르는 아티팩트: {k} (가능: {', '.join(ARTIFACTS)})")

    state = job.load_state() if opts.reuse else {}
    state.setdefault("job_id", job.id)
    state["started_at"] = state.get("started_at") or now()
    state["pack"] = pack.relative_to(ROOT).as_posix()
    state["requested_artifacts"] = kinds
    state["options"] = {
        k: v for k, v in vars(opts).items() if k not in ("func", "source")
    }

    log(f"작업 {job.id} · 팩 {pack.name} · 아티팩트 {', '.join(kinds)}")

    auth_ok, auth_text = ensure_auth(opts.profile, log_fn=log)
    if not auth_ok:
        fail(state, "auth", auth_text[:300])
        job.save_state(state)
        raise SystemExit(
            "NotebookLM에 인증되지 않았다(무인 갱신도 실패). 사용자가 `nlm login` 을 한 번 돌려야 한다.\n"
            "돌리지 못했으면 발표팩까지만 완성된 것이다 — 생성한 척 보고하지 않는다."
        )

    if not state.get("notebook_id"):
        title = opts.notebook_title or f"{job.id}"
        if not step_create_notebook(job, state, title, opts.profile):
            job.save_state(state)
            return 1
        job.save_state(state)
        step_persona(job, state, instruction, opts.profile)
        job.save_state(state)
        if not step_upload(job, state, [pack] + extra_files, opts.url or [], opts.profile):
            fail(state, "upload", "올라간 자료가 하나도 없다. 아티팩트를 만들지 않는다.")
            job.save_state(state)
            write_report(job, state)
            return 1
        job.save_state(state)
    else:
        log(f"  기존 노트북 재사용: {state['notebook_id']}")

    step_create_artifacts(job, state, kinds, opts, instruction)
    job.save_state(state)
    log("  생성 대기 — NotebookLM은 비동기다. 몇 분 걸린다.")
    step_poll(job, state, kinds, opts)
    step_download(job, state, kinds, opts)
    step_inspect(job, state, pack_text)
    state["finished_at"] = now()
    job.save_state(state)
    write_report(job, state)

    done = [k for k, v in state.get("artifacts", {}).items() if v.get("status") == "completed"]
    log(f"완료 {len(done)}/{len(kinds)} — {', '.join(done) or '없음'}")
    flagged = [k for k, r in state.get("inspection", {}).items()
               if not r.get("ok") or r.get("novel_numbers")]
    if flagged:
        log(f"검사에서 걸린 것: {', '.join(flagged)} — 닫기 전에는 promote 가 보류한다.")
    if state.get("errors"):
        log("실패·미완료 항목이 있다. _nlm_run.md 를 본다.")
        return 2
    return 0


def cmd_resume(opts: argparse.Namespace) -> int:
    job = Job(opts.id)
    state = job.load_state()
    if not state.get("notebook_id"):
        raise SystemExit(f"이어받을 실행 기록이 없다: {job.state_path}")
    pack = ROOT / state["pack"]
    pack_text = pack.read_text(encoding="utf-8", errors="replace")
    kinds = state.get("requested_artifacts", DEFAULT_ARTIFACTS)
    state["errors"] = []
    log(f"이어받기: 노트북 {state['notebook_id']}")
    step_poll(job, state, kinds, opts)
    step_download(job, state, kinds, opts)
    step_inspect(job, state, pack_text)
    state["finished_at"] = now()
    job.save_state(state)
    write_report(job, state)
    return 2 if state.get("errors") else 0


PROMOTE_NAMES = {
    ("report", ".md"): "_브리핑.md",
    ("infographic", ".png"): "_인포그래픽.png",
    ("slides", ".pptx"): "_슬라이드.pptx",
    ("slides", ".pdf"): "_슬라이드.pdf",
}


def cmd_promote(opts: argparse.Namespace) -> int:
    job = Job(opts.id)
    state = job.load_state()
    if not state:
        raise SystemExit(f"실행 기록이 없다: {job.state_path}")

    blockers = []
    for key, res in state.get("inspection", {}).items():
        if not res.get("ok"):
            blockers.append(f"{key}: 기계 검사 미통과")
        if res.get("novel_numbers"):
            blockers.append(f"{key}: 팩에 없는 숫자 {len(res['novel_numbers'])}건")
    if blockers and not opts.force:
        print("이관 보류 — 닫지 않은 항목이 있다:")
        for b in blockers:
            print(f"  - {b}")
        print("\n확인해서 닫았으면 --force 를 붙인다. 확인 없이 붙이지 마라.")
        return 1

    final_root = ROOT / "최종"
    dest_dir = final_root / job.id
    dest_dir.mkdir(parents=True, exist_ok=True)
    moved = []
    for kind, info in state.get("artifacts", {}).items():
        for suffix, rel in (info.get("files") or {}).items():
            src = ROOT / rel
            if not src.exists():
                continue
            name = PROMOTE_NAMES.get((kind, suffix), f"_{kind}{suffix}")
            dst = dest_dir / f"{job.id}{name}"
            shutil.copy2(src, dst)
            moved.append(dst.relative_to(ROOT).as_posix())
            print(f"  → {dst.relative_to(ROOT).as_posix()}")

    pack = ROOT / state["pack"]
    if pack.exists():
        dst = dest_dir / f"{job.id}_발표팩.md"
        shutil.copy2(pack, dst)
        moved.append(dst.relative_to(ROOT).as_posix())
        print(f"  → {dst.relative_to(ROOT).as_posix()}")

    if not moved:
        print("이관할 파일이 없다.")
        return 1

    state["promoted"] = {"at": now(), "files": moved, "forced": bool(opts.force)}
    job.save_state(state)
    manifest = final_root / "_manifest.md"
    line = f"- {now()[:10]} · `{job.id}` · " + " · ".join(Path(m).name for m in moved)
    with manifest.open("a", encoding="utf-8") as fh:
        fh.write(line + "\n")
    print(f"\n{len(moved)}개 이관. 목록: 최종/_manifest.md")
    return 0


# ── 인자 ────────────────────────────────────────────────────────────────────


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="nlm_pipeline.py",
        description="발표팩 → NotebookLM → 슬라이드 (기본). --artifacts 로 인포그래픽·브리핑도",
    )
    sub = p.add_subparsers(dest="cmd", required=True)

    c = sub.add_parser("check", help="nlm 설치·인증 확인")
    c.add_argument("--profile")
    c.set_defaults(func=lambda o: step_check(o.profile))

    def common(sp):
        sp.add_argument("--id", required=True, help="작업 ID")
        sp.add_argument("--profile", help="nlm 프로필")
        sp.add_argument("--timeout", type=int, default=2400, help="생성 대기 상한(초)")
        sp.add_argument("--language", default="ko", help="BCP-47 (기본 ko, 끄려면 빈 문자열)")

    r = sub.add_parser("run", help="전체 파이프라인")
    common(r)
    r.add_argument("--pack", help="발표팩 경로 (기본: workspace/07_...pack*.md 중 최신)")
    r.add_argument("--source", action="append", help="추가 업로드 파일 (반복 가능)")
    r.add_argument("--url", action="append", help="추가 업로드 URL (반복 가능)")
    r.add_argument("--artifacts", default=",".join(DEFAULT_ARTIFACTS),
                   help="report,infographic,slides 중 쉼표 구분 (기본: slides만)")
    r.add_argument("--notebook-title")
    r.add_argument("--report-format", default="Briefing Doc",
                   choices=["Briefing Doc", "Study Guide", "Blog Post", "Create Your Own"])
    r.add_argument("--slide-format", default="detailed_deck",
                   choices=["detailed_deck", "presenter_slides"])
    r.add_argument("--slide-length", default="default", choices=["short", "default"])
    r.add_argument("--orientation", default="landscape",
                   choices=["landscape", "portrait", "square"])
    r.add_argument("--detail", default="standard", choices=["concise", "standard", "detailed"])
    r.add_argument("--style", default="professional")
    r.add_argument("--reuse", action="store_true", help="기존 _nlm_run.json 을 이어 쓴다")
    r.add_argument("--redownload", action="store_true", help="이미 받은 파일도 지우고 다시 받는다")
    r.add_argument("--slides-pptx", action="store_true",
                   help="슬라이드를 pptx로도 받는다 (장마다 통 이미지라 글자 수정은 안 된다)")
    r.set_defaults(func=cmd_run)

    s = sub.add_parser("resume", help="폴링·다운로드·검사만 다시")
    common(s)
    s.add_argument("--redownload", action="store_true", help="이미 받은 파일도 지우고 다시 받는다")
    s.add_argument("--slides-pptx", action="store_true", help="슬라이드를 pptx로도 받는다")
    s.set_defaults(func=cmd_resume)

    m = sub.add_parser("promote", help="검사 통과분을 최종/<ID>/ 로")
    m.add_argument("--id", required=True)
    m.add_argument("--force", action="store_true", help="보류 항목을 직접 확인해 닫았을 때만")
    m.set_defaults(func=cmd_promote)

    return p


def main() -> int:
    os.environ.setdefault("PYTHONIOENCODING", "utf-8")
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    opts = build_parser().parse_args()
    return opts.func(opts) or 0


if __name__ == "__main__":
    raise SystemExit(main())
